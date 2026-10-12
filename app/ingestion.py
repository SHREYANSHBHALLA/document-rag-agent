
"""Load PDF, DOCX, and TXT files while preserving document structure."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions
from docling.document_converter import DocumentConverter, PdfFormatOption
from docx import Document as DocxDocument


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100

CHUNK_SEPARATORS: tuple[str, ...] = (
    "\n\n", "\n", ". ", " ", ""
)

REFERENCE_HEADING = re.compile(
    r"^(?:(?:[IVXLCDM]+|\d+)\.?\s+)?"
    r"(?:references|bibliography|works\s+cited|literature\s+cited)\s*:?\s*$",
    re.IGNORECASE,
)

ABSTRACT_HEADING = re.compile(
    r"^\s*abstract\s*:?\s*$",
    re.IGNORECASE,
)

ABSTRACT_PREFIX = re.compile(
    r"^\s*abstract\s*[-–—:]\s*",
    re.IGNORECASE,
)

INDEX_TERMS_PREFIX = re.compile(
    r"^\s*index\s+terms?\s*[-–—:]",
    re.IGNORECASE,
)

NUMBERED_SECTION_HEADING = re.compile(
    r"^(?:[IVXLCDM]+|[A-Z]|\d+(?:\.\d+)*)\.\s+\S+",
    re.IGNORECASE,
)

COMMON_BODY_HEADINGS = {
    "introduction",
    "overview",
    "background",
    "related work",
    "methodology",
    "methods",
    "discussion",
    "conclusion",
}


@dataclass(frozen=True)
class Document:
    """Text with its source, page, and structural metadata."""

    text: str
    source: str
    page: int | None
    section_type: str = "body"
    section_title: str | None = None


def _label(item: object) -> str:
    value = getattr(item, "label", "text")
    return str(getattr(value, "value", value)).lower()


def _page_number(item: object) -> int | None:
    for provenance in getattr(item, "prov", None) or []:
        page = getattr(provenance, "page_no", None)
        if page is not None:
            return int(page)
    return None


def _is_reference_heading(text: str) -> bool:
    text = re.sub(r"^\s*#+\s*", "", text.strip())
    return bool(REFERENCE_HEADING.fullmatch(text))


def _is_abstract_body(text: str) -> bool:
    return bool(ABSTRACT_PREFIX.match(text))


def _export_item_text(item: object, parsed_document: object) -> str:
    """Export structured items when possible, with safe fallbacks."""

    label = _label(item)

    if label == "table":
        exporter = getattr(item, "export_to_markdown", None)
        if callable(exporter):
            try:
                return str(exporter(doc=parsed_document) or "").strip()
            except Exception:
                pass

    return str(getattr(item, "text", "") or "").strip()


def _group_elements(
    elements: Sequence[Document],
) -> list[Document]:
    """Combine adjacent elements with matching metadata."""

    grouped: list[Document] = []

    for item in elements:
        if (
            grouped
            and grouped[-1].source == item.source
            and grouped[-1].page == item.page
            and grouped[-1].section_type == item.section_type
            and grouped[-1].section_title == item.section_title
        ):
            previous = grouped[-1]
            grouped[-1] = Document(
                text=f"{previous.text}\n\n{item.text}",
                source=item.source,
                page=item.page,
                section_type=item.section_type,
                section_title=item.section_title,
            )
        else:
            grouped.append(item)

    return grouped


def load_document(file_path: str) -> list[Document]:
    """Load a supported file and preserve available structure."""

    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {file_path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported extension '{extension}'. "
            f"Supported types: {supported}."
        )

    if extension == ".pdf":
        return _load_pdf(path, path.name)
    if extension == ".docx":
        return _load_docx(path, path.name)

    return _load_txt(path, path.name)


def _load_pdf(path: Path, source: str) -> list[Document]:
    """Extract Docling items without discarding their labels."""

    options = PdfPipelineOptions()

    converter = DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(
                pipeline_options=options
            )
        }
    )

    parsed = converter.convert(str(path)).document

    items: list[tuple[object, str, str]] = []

    for item, _level in parsed.iterate_items():
        label = _label(item)

        if label in {"page_header", "page_footer", "marker", "empty_value"}:
            continue

        text = _export_item_text(item, parsed)
        if text:
            items.append((item, label, text))

    documents: list[Document] = []
    current_type = "front_matter"
    current_title: str | None = None
    first_section_header_seen = False
    references_started = False

    for item, label, text in items:
        page = _page_number(item)

        # References are detected from the heading or Docling's item label.
        if _is_reference_heading(text) or label == "reference":
            references_started = True
            current_type = "references"
            current_title = "References"
            section_type = "references"

        elif references_started:
            section_type = "references"

        elif label == "title":
            current_type = "front_matter"
            current_title = text
            section_type = "front_matter"

        elif ABSTRACT_HEADING.fullmatch(text) or _is_abstract_body(text):
            current_type = "abstract"
            current_title = "Abstract"
            section_type = "abstract"

        elif INDEX_TERMS_PREFIX.match(text):
            current_type = "front_matter"
            current_title = "Index Terms"
            section_type = "front_matter"

        elif label == "section_header":
            normalized_heading = text.strip().casefold()

            looks_like_title = (
                not first_section_header_seen
                and page == 1
                and not NUMBERED_SECTION_HEADING.match(text)
                and normalized_heading not in COMMON_BODY_HEADINGS
                and len(text.strip()) > 12
            )

            first_section_header_seen = True

            if looks_like_title:
                current_type = "front_matter"
                current_title = text
                section_type = "front_matter"
            else:
                current_type = "body"
                current_title = text
                section_type = "body"

        elif label == "table":
            section_type = "table"

        elif label == "caption":
            section_type = "figure_caption"

        elif label in {"picture", "chart"}:
            section_type = "figure_caption"

        elif label == "footnote":
            section_type = "footnote"

        else:
            # Author names and affiliations inherit front_matter until
            # a recognized abstract or body heading changes the state.
            section_type = current_type

        documents.append(
            Document(
                text=text,
                source=source,
                page=page,
                section_type=section_type,
                section_title=current_title,
            )
        )

    return _group_elements(documents)


def _load_docx(path: Path, source: str) -> list[Document]:
    """Extract DOCX paragraphs and use heading styles where available."""

    docx = DocxDocument(path)
    documents: list[Document] = []

    current_type = "body"
    current_title: str | None = None
    references_started = False

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = (getattr(paragraph.style, "name", "") or "").lower()

        if _is_reference_heading(text):
            references_started = True
            current_type = "references"
            current_title = "References"
            section_type = "references"

        elif references_started:
            section_type = "references"

        elif style in {"title", "subtitle"}:
            current_type = "front_matter"
            current_title = text
            section_type = "front_matter"

        elif ABSTRACT_HEADING.fullmatch(text) or _is_abstract_body(text):
            current_type = "abstract"
            current_title = "Abstract"
            section_type = "abstract"

        elif INDEX_TERMS_PREFIX.match(text):
            current_type = "front_matter"
            current_title = "Index Terms"
            section_type = "front_matter"

        elif style.startswith("heading"):
            current_type = "body"
            current_title = text
            section_type = "body"

        else:
            section_type = current_type

        documents.append(
            Document(
                text=text,
                source=source,
                page=None,
                section_type=section_type,
                section_title=current_title,
            )
        )

    return _group_elements(documents)


def _load_txt(path: Path, source: str) -> list[Document]:
    """Load text files with simple heading-based classification."""

    text = path.read_text(
        encoding="utf-8",
        errors="replace",
    ).strip()

    if not text:
        return []

    documents: list[Document] = []
    current_type = "body"
    current_title: str | None = None
    references_started = False
    buffer: list[str] = []

    def flush() -> None:
        content = "\n".join(buffer).strip()
        if content:
            documents.append(
                Document(
                    text=content,
                    source=source,
                    page=None,
                    section_type=current_type,
                    section_title=current_title,
                )
            )
        buffer.clear()

    for line in text.splitlines():
        cleaned = line.strip()

        if _is_reference_heading(cleaned):
            flush()
            references_started = True
            current_type = "references"
            current_title = "References"
            buffer.append(line)

        elif references_started:
            buffer.append(line)

        elif ABSTRACT_HEADING.fullmatch(cleaned) or _is_abstract_body(cleaned):
            flush()
            current_type = "abstract"
            current_title = "Abstract"
            buffer.append(line)

        elif INDEX_TERMS_PREFIX.match(cleaned):
            flush()
            current_type = "front_matter"
            current_title = "Index Terms"
            buffer.append(line)

        elif cleaned.startswith("#"):
            flush()
            current_type = "body"
            current_title = cleaned.lstrip("#").strip()
            buffer.append(line)

        else:
            buffer.append(line)

    flush()
    return _group_elements(documents)


def chunk_documents(
    documents: Sequence[Document],
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
) -> list[Document]:
    """Recursively chunk text while preserving source and section metadata."""

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than 0.")
    if chunk_overlap < 0:
        raise ValueError("chunk_overlap must be 0 or greater.")
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size.")

    chunks: list[Document] = []

    for document in _group_elements(documents):
        for text in _chunk_text(
            document.text,
            chunk_size,
            chunk_overlap,
        ):
            chunks.append(
                Document(
                    text=text,
                    source=document.source,
                    page=document.page,
                    section_type=document.section_type,
                    section_title=document.section_title,
                )
            )

    return chunks


def _chunk_text(
    text: str,
    chunk_size: int,
    chunk_overlap: int,
) -> list[str]:
    pieces = _split_recursively(
        text.strip(),
        chunk_size,
        CHUNK_SEPARATORS,
    )

    return _merge_with_overlap(
        pieces,
        chunk_size,
        chunk_overlap,
    )


def _split_recursively(
    text: str,
    chunk_size: int,
    separators: Sequence[str],
) -> list[str]:
    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    separator, *remaining = separators or ("",)
    parts = list(text) if separator == "" else text.split(separator)

    results: list[str] = []
    buffer: list[str] = []
    buffer_length = 0

    for part in parts:
        extra = len(separator) if buffer else 0

        if buffer_length + extra + len(part) <= chunk_size:
            buffer.append(part)
            buffer_length += extra + len(part)
            continue

        if buffer:
            results.append(separator.join(buffer))
            buffer = []
            buffer_length = 0

        if len(part) > chunk_size:
            results.extend(
                _split_recursively(part, chunk_size, remaining)
            )
        else:
            buffer = [part]
            buffer_length = len(part)

    if buffer:
        results.append(separator.join(buffer))

    return [piece.strip() for piece in results if piece.strip()]


def _merge_with_overlap(
    pieces: Sequence[str],
    chunk_size: int,
    chunk_overlap: int,
) -> list[str]:
    if not pieces:
        return []

    merged: list[str] = []
    current = pieces[0]

    for piece in pieces[1:]:
        candidate = f"{current}\n{piece}"

        if len(candidate) <= chunk_size:
            current = candidate
            continue

        merged.append(current)

        if chunk_overlap:
            overlap_text = current[-chunk_overlap:].lstrip()
            candidate = (
                f"{overlap_text}\n{piece}"
                if overlap_text
                else piece
            )
            current = candidate if len(candidate) <= chunk_size else piece
        else:
            current = piece

    merged.append(current)
    return merged